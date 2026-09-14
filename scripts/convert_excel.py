from __future__ import annotations

import argparse
import json
import re
from datetime import date, datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = Path(r"C:\Users\h2377\Downloads\大学生竞赛小程序_首批30场赛事数据底稿_v1.0.xlsx")
DEFAULT_OUTPUT = ROOT / "data" / "seed" / "competitions.json"

PROFESSIONS = [
    "计算机与AI", "电子信息", "机械与自动化", "数学与统计", "经管与商科",
    "设计与艺术", "建筑与土木", "材料与化工", "能源与环境", "生物与医药",
    "农林与食品", "人文与法学", "外语与教育", "综合/不限专业",
]

PROFESSION_MAP = {
    "计算机与AI": ["计算机与AI"],
    "电子信息": ["电子信息"],
    "机械与自动化": ["机械与自动化"],
    "数理统计": ["数学与统计"],
    "经管与商科": ["经管与商科"],
    "设计与传媒": ["设计与艺术", "人文与法学"],
    "土木建筑": ["建筑与土木"],
    "化学化工与能源": ["材料与化工", "能源与环境"],
    "生命科学与医学": ["生物与医药"],
    "物流与交通": ["经管与商科"],
    "外语与人文": ["外语与教育", "人文与法学"],
    "综合创新创业": ["综合/不限专业"],
}

RECOMMENDATIONS = {"强烈推荐", "值得参加", "适合练手", "高难挑战", "谨慎报名", "信息不足"}
STATUS_VALUES = {"时间待公布", "即将报名", "报名中", "已截止"}
PARTICIPATION_FORMATS = {"个人赛", "团队赛", "个人或团队均可", "暂未明确"}
TEAM_SIZE_CATEGORIES = {"1人", "2～3人", "4～5人", "6人及以上", "人数范围不固定", "暂未明确"}
FEE_TYPES = {"免费", "收费", "不同赛项费用不同", "暂未公布"}
REGISTRATION_ORGANIZATIONS = {
    "自行报名", "学校统一报名", "两种方式均可", "需先校内选拔", "以赛区或学校通知为准", "暂未明确",
}


def iso(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        return value.isoformat()
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=timezone.utc).isoformat()
    return str(value).strip() or None


def slug_id(raw: str) -> str:
    return raw.lower().replace("_", "-")


def split_values(raw) -> list[str]:
    if not raw:
        return []
    return [item.strip() for item in re.split(r"[、，,；;]", str(raw)) if item.strip()]


def normalize_professions(raw) -> list[str]:
    result: list[str] = []
    for old in split_values(raw):
        for new in PROFESSION_MAP.get(old, ["综合/不限专业"]):
            if new not in result:
                result.append(new)
    return result or ["综合/不限专业"]


def normalize_level(raw) -> tuple[str, list[str]]:
    text = str(raw or "").strip()
    if text.startswith("全国性赛事"):
        tags = ["全国性赛事"]
        if "国际赛道" in text:
            tags.append("含国际赛道")
        if "国际中文组" in text:
            tags.append("含国际中文组")
        return "国家级", tags
    if "省" in text or "区域" in text:
        return "省级/区域级", [text] if text else []
    return "其他/未明确", [text] if text else []


def normalize_participation_format(raw) -> str:
    """只按底稿明确措辞归类，不从赛事名称或常识推断。"""
    text = str(raw or "").strip()
    if not text or "为主" in text or "按作品类别" in text or "按赛道" in text or "按组别" in text:
        return "暂未明确"
    if any(token in text for token in ("个人或团队", "个人/团队", "个人和团队")):
        return "个人或团队均可"
    has_personal = "个人" in text or text.startswith("1名") or text.startswith("1人")
    has_team = any(token in text for token in ("团队", "组队", "组成一队", "每队", "战队"))
    if has_personal and has_team:
        return "个人或团队均可"
    if has_team:
        return "团队赛"
    if has_personal:
        return "个人赛"
    return "暂未明确"


def normalize_team_size(raw) -> str:
    """从明确的人数字样分桶；范围或口径不清时保留为不固定/未明确。"""
    text = str(raw or "").strip()
    if not text:
        return "暂未明确"
    first_clause = re.split(r"[；;]", text)[0]
    if any(token in first_clause for token in ("按作品类别", "按组别", "按赛道", "按规则", "人数和", "人数及")):
        return "人数范围不固定"
    numbers = [int(value) for value in re.findall(r"\d+", first_clause)]
    if not numbers:
        return "人数范围不固定" if any(token in first_clause for token in ("团队", "组队", "战队", "个人或")) else "暂未明确"
    maximum = max(numbers)
    if maximum == 1 and min(numbers) == 1:
        return "1人"
    if maximum <= 3:
        return "2～3人"
    if maximum <= 5:
        return "4～5人"
    return "6人及以上"


def normalize_fee_type(raw) -> str:
    text = str(raw or "").strip()
    if not text or any(token in text for token in ("待", "未公布", "未明确", "确认")):
        return "暂未公布"
    if "不同" in text and any(token in text for token in ("赛项", "组别", "地区")):
        return "不同赛项费用不同"
    if any(token in text for token in ("免费", "不收费", "免收")):
        return "免费"
    if any(token in text for token in ("收费", "缴费", "元")):
        return "收费"
    return "暂未公布"


def normalize_registration_organization(raw) -> str:
    """报名组织方式独立于参赛形式，仅使用报名方式原文中的明确组织动作。"""
    text = str(raw or "").strip()
    if not text:
        return "暂未明确"
    if any(token in text for token in ("校内选拔", "学校先组织初赛", "由学校选拔", "组织遴选", "学校组织校赛")):
        return "需先校内选拔"
    if any(token in text for token in ("两种方式", "均可报名", "个人及学校", "个人或学校")):
        return "两种方式均可"
    if any(token in text for token in ("由学校或赛区", "学校和赛区组织", "按学校组织要求", "赛区统一报名", "以赛区", "以学校通知")):
        return "以赛区或学校通知为准"
    if any(token in text for token in ("由学校统一", "学校统一", "学校管理员分配", "学校集体", "由学校组织", "院系先获得")):
        return "学校统一报名"
    if any(token in text for token in ("登录", "官网报名", "在线报名", "自行报名", "个人在", "团队在", "注册")):
        return "自行报名"
    return "暂未明确"


def extract_year(name: str) -> int | None:
    years = [int(value) for value in re.findall(r"20\d{2}", name or "")]
    return max(years) if years else None


def infer_notice_type(url: str, homepage: str | None) -> tuple[str, bool]:
    host = urlparse(url).netloc.lower()
    home_host = urlparse(homepage or "").netloc.lower()
    if home_host and (host == home_host or host.endswith("." + home_host)):
        return "organizer_notice", True
    if host.endswith("gov.cn"):
        return "authority_notice", True
    if host.endswith("edu.cn"):
        return "university_notice", False
    if "xinhuanet.com" in host or "people.com.cn" in host:
        return "media_report", False
    return "third_party_lead", False


def source_record(source_id: str, edition_id: str, url: str | None, source_type: str, official: bool, last_verified_at: str | None):
    return {
        "id": source_id,
        "edition_id": edition_id,
        "url": url,
        "source_type": source_type,
        "is_official": official,
        "published_at": None,
        "last_verified_at": last_verified_at,
        "validity_status": "unchecked",
    }


def read_rows(workbook, sheet_name: str):
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[3]]
    for row in sheet.iter_rows(min_row=4, values_only=True):
        if not row[0]:
            continue
        yield {headers[index]: row[index] for index in range(len(headers))}


def convert(input_path: Path, output_path: Path) -> dict:
    workbook = load_workbook(input_path, data_only=False, read_only=True)
    master_rows = list(read_rows(workbook, "赛事主体"))
    edition_rows = list(read_rows(workbook, "赛事届次"))
    masters_by_legacy_id = {str(row["赛事ID"]): row for row in master_rows}

    output = {
        "meta": {
            "schema_version": "1.0.0",
            "generated_at": "2026-08-20T00:00:00+08:00",
            "source_file": input_path.name,
            "source_sheet_count": len(workbook.sheetnames),
            "source_sheets": workbook.sheetnames,
            "record_count": len(edition_rows),
            "development_preview_only": True,
        },
        "profession_categories": [
            {"id": f"profession-{index:02d}", "name": name, "parent_id": None, "level": "category", "aliases": []}
            for index, name in enumerate(PROFESSIONS, start=1)
        ],
        "event_masters": [],
        "event_editions": [],
        "event_tracks": [],
        "event_schedules": [],
        "data_sources": [],
        "event_professions": [],
        "recommendations": [],
        "migration_issues": [],
    }
    profession_ids = {item["name"]: item["id"] for item in output["profession_categories"]}

    for row in master_rows:
        legacy_id = str(row["赛事ID"])
        event_id = f"event-{slug_id(legacy_id)}"
        level, level_tags = normalize_level(row["赛事级别"])
        professions = normalize_professions(row["适合专业大类"])
        aliases = split_values(row["别名/旧名"])
        output["event_masters"].append({
            "id": event_id,
            "legacy_id": legacy_id,
            "standard_name": str(row["赛事主体"]).strip(),
            "aliases": aliases,
            "website_url": iso(row["官方主页"]),
            "long_term_organizer": str(row["赛事主体组织单位"] or "").strip(),
            "description": f"面向{str(row['典型参赛对象'] or '高校学生').strip()}的竞赛。",
            "event_level": level,
            "objective_tags": list(dict.fromkeys(level_tags + split_values(row["客观标签"]))),
            "profession_categories": professions,
            "last_verified_at": iso(row["最后核验日期"]),
            "migration": {"source_status": row["主体资料状态"], "legacy_primary_profession": row["主专业大类"]},
        })

    for row in edition_rows:
        legacy_event_id = str(row["赛事ID"])
        master = masters_by_legacy_id[legacy_event_id]
        event_id = f"event-{slug_id(legacy_event_id)}"
        year = extract_year(str(row["当前届次"] or ""))
        suffix = str(year) if year else "undated"
        edition_id = f"edition-{slug_id(legacy_event_id)}-{suffix}"
        track_id = f"track-{slug_id(legacy_event_id)}-{suffix}-01"
        schedule_id = f"schedule-{slug_id(legacy_event_id)}-{suffix}-01"
        level, level_tags = normalize_level(master["赛事级别"])
        professions = normalize_professions(master["适合专业大类"])
        start = iso(row["报名开始"])
        end = iso(row["报名截止"])
        status_override = row["状态覆盖"] if row["状态覆盖"] in STATUS_VALUES else None
        override_reason = str(row["时间说明"] or row["管理员备注"] or "").strip() if status_override else None
        recommendation_value = str(row["平台推荐等级（仅供参考）"] or "信息不足").strip()
        if recommendation_value not in RECOMMENDATIONS:
            recommendation_value = "信息不足"
        raw_reasons = [part.strip() for part in re.split(r"[；;]", str(row["推荐理由"] or "")) if part.strip()]
        reasons = raw_reasons[:3] or ["当前结构化资料不足，暂不作正负面判断"]
        issues = ["source_reachability_not_checked"]
        track_name = str(row["赛项/赛道"] or "未细分赛项").strip()
        if any(token in track_name for token in ("/", "、", "等", "主体赛及", "类别")):
            issues.append("track_requires_official_split")
            output["migration_issues"].append({"edition_id": edition_id, "type": "track_split_needed", "detail": track_name})
        if year is None:
            issues.append("edition_year_missing")
        if not start:
            issues.append("registration_start_pending")
        if not end:
            issues.append("registration_end_pending")

        raw_team_size = str(row["组队人数"] or "").strip()
        raw_registration_method = str(row["官方报名方式"] or "").strip()
        raw_fees = "待官方资料确认"
        output["event_editions"].append({
            "id": edition_id,
            "legacy_id": row["届次ID"],
            "event_id": event_id,
            "edition_name": str(row["当前届次"] or "").strip(),
            "year": year,
            "event_level": level,
            "objective_tags": level_tags,
            "organizer": str(row["主办单位"] or "").strip(),
            "participants": str(row["参赛对象"] or "").strip(),
            "eligibility": str(row["参赛对象"] or "").strip(),
            "format_and_rules": str(row["基本比赛规则"] or "").strip(),
            "team_size": raw_team_size,
            "participation_format": normalize_participation_format(raw_team_size),
            "team_size_category": normalize_team_size(raw_team_size),
            "registration_method": raw_registration_method,
            "registration_organization": normalize_registration_organization(raw_registration_method),
            "registration_url": iso(row["官方报名入口"]),
            "fees": raw_fees,
            "fee_type": normalize_fee_type(raw_fees),
            "data_status": "pending_manual_review",
            "system_validation": "not_run",
            "manual_review": "pending",
            "data_version": str(row["数据版本"] or "v1.0"),
            "last_verified_at": iso(row["最后核验日期"]),
            "quality_issues": issues,
            "development_preview": True,
            "migration": {"excel_system_validation": row["系统验证"], "excel_manual_review": row["人工审核"], "excel_admin_note": row["管理员备注"]},
        })

        output["event_tracks"].append({
            "id": track_id,
            "edition_id": edition_id,
            "name": track_name,
            "scope_note": "组合赛项需按后续官方资料拆分" if "track_requires_official_split" in issues else None,
        })
        output["event_schedules"].append({
            "id": schedule_id,
            "edition_id": edition_id,
            "track_id": track_id,
            "registration_start": start,
            "registration_end": end,
            "competition_start": iso(row["比赛开始"]),
            "competition_end": iso(row["比赛结束"]),
            "pending_announcement": not bool(start),
            "status_override": status_override,
            "status_override_reason": override_reason,
            "time_note": str(row["时间说明"] or "").strip(),
            "change_status": "normal",
        })

        last_verified = iso(row["最后核验日期"])
        homepage = iso(master["官方主页"])
        output["data_sources"].append(source_record(f"source-{slug_id(legacy_event_id)}-{suffix}-homepage", edition_id, homepage, "organizer_website", True, last_verified))
        registration_url = iso(row["官方报名入口"])
        if registration_url:
            output["data_sources"].append(source_record(f"source-{slug_id(legacy_event_id)}-{suffix}-registration", edition_id, registration_url, "official_registration_system", True, last_verified))
        notice_url = iso(row["官方通知来源"])
        if notice_url:
            notice_type, official = infer_notice_type(notice_url, homepage)
            notice_id = f"source-{slug_id(legacy_event_id)}-{suffix}-notice"
            output["data_sources"].append(source_record(notice_id, edition_id, notice_url, notice_type, official, last_verified))
        else:
            notice_id = f"source-{slug_id(legacy_event_id)}-{suffix}-homepage"

        for profession in professions:
            output["event_professions"].append({
                "id": f"event-profession-{slug_id(legacy_event_id)}-{suffix}-{profession_ids[profession].split('-')[-1]}",
                "edition_id": edition_id,
                "profession_id": profession_ids[profession],
            })

        evidence_refs = [f"event_edition:{edition_id}:organizer", f"event_track:{track_id}:name", f"data_source:{notice_id}"][:len(reasons)]
        output["recommendations"].append({
            "id": f"recommendation-{slug_id(legacy_event_id)}-{suffix}-v1",
            "edition_id": edition_id,
            "recommendation": recommendation_value,
            "suitable_for": f"{ '、'.join(professions) }相关方向；{str(row['参赛对象'] or '具体资格以官方通知为准').strip()}",
            "reasons": reasons,
            "evidence_refs": evidence_refs,
            "limitations": [str(row["管理员备注"] or "具体投入、资格和校内认定以官方最新通知及本校文件为准").strip()][:3],
            "preparation": "先核对当届官方通知、报名路径和赛项要求，再按截止时间倒排准备。",
            "confidence": "中" if issues else "高",
            "missing_information": [issue for issue in issues if issue.endswith("pending")],
            "event_data_version": str(row["数据版本"] or "v1.0"),
            "rule_version": "prd-v1.1",
            "system_validation": "not_run",
            "manual_review": "pending",
            "published_cache": False,
        })

        legacy_year = extract_year(str(row["届次ID"] or ""))
        if year and legacy_year and year != legacy_year:
            output["migration_issues"].append({"edition_id": edition_id, "type": "legacy_id_year_corrected", "detail": f"{row['届次ID']} -> {edition_id}"})

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    if not args.input.exists():
        raise SystemExit(f"Excel not found: {args.input}")
    data = convert(args.input, args.output)
    print(f"converted editions={len(data['event_editions'])} sources={len(data['data_sources'])} output={args.output}")


if __name__ == "__main__":
    main()
